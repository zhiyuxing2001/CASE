"""病案病程（系列）列表。

一个「病案」是一段病程：初诊为系列头（father_id 指向自身），后续复诊/
随诊同属该病案（course_id 相同）。列表按 course_id 聚合，每病案一行。
同一患者可有多个病案，因此病案标识用独立的 course_id，而非 patient_id。
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from .. import schemas
from ..services.report_service import generate_case_report
from .deps import get_db

router = APIRouter(prefix="/api/courses", tags=["courses"])

_SELECT = """
SELECT
  r.course_id    AS course_id,
  r.record_id    AS first_record_id,
  r.patient_id   AS patient_id,
  p.patient_name AS patient_name,
  r.clinic_date  AS first_date,
  n.complaint    AS complaint,
  d.syndrome     AS syndrome,
  d.tcm_disease  AS tcm_disease,
  m.mentor_name  AS mentor_name,
  (SELECT COUNT(*) FROM info_record v
     WHERE v.course_id = r.course_id AND v.is_deleted = 0) AS visit_count,
  (SELECT MAX(v.clinic_date) FROM info_record v
     WHERE v.course_id = r.course_id AND v.is_deleted = 0) AS last_date,
  EXISTS(SELECT 1 FROM info_record v
           JOIN prescription_item pi ON pi.record_id = v.record_id
          WHERE v.course_id = r.course_id AND pi.needs_review = 1) AS needs_review
FROM info_record r
JOIN info_patient p ON p.patient_id = r.patient_id
LEFT JOIN case_narrative n ON n.record_id = r.record_id
LEFT JOIN diagnosis d ON d.record_id = r.record_id
LEFT JOIN mentor m ON m.mentor_id = r.mentor_id
WHERE r.father_id = r.record_id AND r.is_deleted = 0
"""

_COUNT = """
SELECT COUNT(*)
FROM info_record r
JOIN info_patient p ON p.patient_id = r.patient_id
LEFT JOIN case_narrative n ON n.record_id = r.record_id
LEFT JOIN diagnosis d ON d.record_id = r.record_id
WHERE r.father_id = r.record_id AND r.is_deleted = 0
"""


def _as_date(value) -> date | None:
    return date.fromisoformat(str(value)) if value is not None else None


@router.get("", response_model=schemas.CourseList)
def list_courses(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    q: str = "",
    syndrome: str = "",
    mentor_id: str = "",
    db: Session = Depends(get_db),
) -> schemas.CourseList:
    where = ""
    params: dict = {}
    if q:
        where += " AND (p.patient_name LIKE :q OR n.complaint LIKE :q OR d.syndrome LIKE :q)"
        params["q"] = f"%{q}%"
    if syndrome:
        where += " AND d.syndrome LIKE :syndrome"
        params["syndrome"] = f"%{syndrome}%"
    if mentor_id:
        where += " AND r.mentor_id = :mentor_id"
        params["mentor_id"] = mentor_id

    total = db.execute(text(_COUNT + where), params).scalar() or 0
    rows = db.execute(
        text(_SELECT + where + " ORDER BY r.clinic_date DESC LIMIT :lim OFFSET :off"),
        {**params, "lim": page_size, "off": (page - 1) * page_size},
    ).all()

    items = [
        schemas.CourseSummary(
            course_id=row.course_id,
            first_record_id=row.first_record_id,
            patient_id=row.patient_id,
            patient_name=row.patient_name or "",
            first_date=_as_date(row.first_date),
            last_date=_as_date(row.last_date) or _as_date(row.first_date),
            visit_count=row.visit_count or 1,
            complaint=row.complaint or "",
            syndrome=row.syndrome or "",
            tcm_disease=row.tcm_disease or "",
            mentor_name=row.mentor_name or "",
            needs_review=bool(row.needs_review),
        )
        for row in rows
    ]
    return schemas.CourseList(total=total, page=page, page_size=page_size, items=items)


def _in_placeholders(ids: list[str], offset: int = 0) -> str:
    return ", ".join(f":c{i + offset}" for i in range(len(ids)))


def _in_params(ids: list[str], offset: int = 0) -> dict:
    return {f"c{i + offset}": cid for i, cid in enumerate(ids)}


@router.post("/delete", response_model=schemas.CourseBatchResult)
def delete_courses(
    payload: schemas.CourseBatchRequest,
    db: Session = Depends(get_db),
) -> schemas.CourseBatchResult:
    """软删除所选病案：把该病案下所有就诊标记为已删除。"""
    ids = payload.course_ids
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要删除的病案")
    result = db.execute(
        text(f"UPDATE info_record SET is_deleted = 1 WHERE course_id IN ({_in_placeholders(ids)})"),
        _in_params(ids),
    )
    db.commit()
    return schemas.CourseBatchResult(deleted=result.rowcount or 0)


@router.post("/export")
def export_courses(
    payload: schemas.CourseBatchRequest,
    db: Session = Depends(get_db),
) -> StreamingResponse:
    """把所选病案导出为 Excel（病案列表 + 就诊明细两个工作表）。"""
    ids = payload.course_ids
    if not ids:
        raise HTTPException(status_code=400, detail="请选择要导出的病案")

    wb = Workbook()

    # 工作表一：病案列表（一个病案一行）
    ws1 = wb.active
    ws1.title = "病案列表"
    ws1.append(["病案号", "患者", "初诊主诉", "证型", "中医病名",
                "就诊次数", "初诊日期", "最近就诊", "带教老师"])
    summary_rows = db.execute(
        text(_SELECT + f" AND r.course_id IN ({_in_placeholders(ids, 100)})"
             + " ORDER BY r.clinic_date DESC"),
        _in_params(ids, 100),
    ).all()
    for row in summary_rows:
        ws1.append([
            row.course_id, row.patient_name, row.complaint, row.syndrome,
            row.tcm_disease, row.visit_count,
            str(row.first_date), str(row.last_date or row.first_date),
            row.mentor_name,
        ])

    # 工作表二：就诊明细（一次就诊一行，药味合并为文本）
    ws2 = wb.create_sheet("就诊明细")
    ws2.append(["病案号", "患者", "就诊日期", "第几诊", "主诉", "现病史",
                "舌质", "舌苔", "脉象", "中医病名", "证型", "西医诊断",
                "辨证分析", "治法", "方剂名", "付数", "煎煮法", "药味"])
    detail_rows = db.execute(
        text("""
            SELECT r.record_id, r.course_id, p.patient_name, r.clinic_date,
                   r.visit_no, n.complaint, n.present_illness, n.body_of_tongue,
                   n.fur_of_tongue, n.pulse, d.tcm_disease, d.syndrome,
                   d.wm_diagnosis, d.patterns_analysis, t.treatment_principle,
                   t.formula_name, t.dose_count, t.decoction
              FROM info_record r
              JOIN info_patient p ON p.patient_id = r.patient_id
              LEFT JOIN case_narrative n ON n.record_id = r.record_id
              LEFT JOIN diagnosis d ON d.record_id = r.record_id
              LEFT JOIN treatment t ON t.record_id = r.record_id
             WHERE r.course_id IN (""" + _in_placeholders(ids, 200) + """)
               AND r.is_deleted = 0
             ORDER BY r.course_id, r.visit_no
        """),
        _in_params(ids, 200),
    ).all()

    record_ids = [row.record_id for row in detail_rows]
    herbs_by_record: dict[int, str] = {}
    if record_ids:
        herb_rows = db.execute(
            text(f"""
                SELECT record_id, herb_name_norm, dose, unit, decoction_note
                  FROM prescription_item
                 WHERE record_id IN ({_in_placeholders([str(r) for r in record_ids], 300)})
                 ORDER BY sequence
            """),
            _in_params([str(r) for r in record_ids], 300),
        ).all()
        for rid, name, dose, unit, note in herb_rows:
            if dose is not None:
                dose_str = str(int(dose)) if float(dose).is_integer() else str(dose)
                piece = f"{name} {dose_str}{unit}"
            else:
                piece = str(name)
            if note:
                piece += note  # 特殊煎服法，Excel 无上标，直接拼接
            herbs_by_record.setdefault(int(rid), []).append(piece)

    for row in detail_rows:
        ws2.append([
            row.course_id, row.patient_name, str(row.clinic_date), row.visit_no,
            row.complaint, row.present_illness, row.body_of_tongue,
            row.fur_of_tongue, row.pulse, row.tcm_disease, row.syndrome,
            row.wm_diagnosis, row.patterns_analysis, row.treatment_principle,
            row.formula_name, row.dose_count, row.decoction,
            "、".join(herbs_by_record.get(row.record_id, [])),
        ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"case-export-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{course_id}/report")
def course_report(
    course_id: str,
    db: Session = Depends(get_db),
) -> StreamingResponse:
    """把单个病案按模版导出为 Word 报告（含关联学习心得）。"""
    buf = generate_case_report(db, course_id)
    if buf is None:
        raise HTTPException(status_code=404, detail="病案不存在")
    filename = f"case-{course_id[:8]}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
