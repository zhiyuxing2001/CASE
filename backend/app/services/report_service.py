"""病案报告（Word）生成服务。

按模版把一段病程（初诊 + 随诊）连同关联的学习心得导出为 .docx：
标题 → 基本信息 → 病程记录（逐诊）→ 学习心得（含导师点评）。
"""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (CaseNarrative, Diagnosis, ExamReport, InfoPatient,
                      InfoRecord, LabResult, LearningNote, MentorComment,
                      PrescriptionItem, Treatment)

NOTE_TYPES = {0: "跟诊日志", 1: "学习心得", 2: "读书笔记", 3: "病例讨论", 4: "阶段总结"}
ABNORMAL = {0: "", 1: "↑", 2: "↓"}


def _style_run(run, size: int = 11, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_kv(doc: Document, label: str, value) -> None:
    if value in (None, ""):
        return
    p = doc.add_paragraph()
    _style_run(p.add_run(label), bold=True)
    _style_run(p.add_run(str(value)))


def _fmt_dose(dose) -> str:
    if dose is None:
        return ""
    if float(dose).is_integer():
        return str(int(dose))
    return str(dose)


def _add_prescription(doc: Document, herbs) -> None:
    """处方笺：每行四味药，药名+剂量，特殊煎服法写为上角标。"""
    table = doc.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    cells = None
    for i, h in enumerate(herbs):
        if i % 4 == 0:
            cells = table.add_row().cells
        p = cells[i % 4].paragraphs[0]
        name = h.herb_name_norm or h.herb_name
        _style_run(p.add_run(name + _fmt_dose(h.dose) + (h.unit or "g")))
        if h.decoction_note:
            sup = p.add_run(h.decoction_note)
            _style_run(sup)
            sup.font.superscript = True


def _md_to_plain(markdown: str) -> list[str]:
    lines: list[str] = []
    for line in markdown.splitlines():
        line = line.strip()
        if not line:
            continue
        line = re.sub(r"^#+\s*", "", line)
        line = re.sub(r"^\s*[-*+]\s+", "· ", line)
        line = line.replace("**", "").replace("`", "")
        lines.append(line)
    return lines


def _herbs_by_record(db: Session, record_ids: list[int]) -> dict[int, list]:
    rows = db.execute(
        select(PrescriptionItem)
        .where(PrescriptionItem.record_id.in_(record_ids))
        .order_by(PrescriptionItem.record_id, PrescriptionItem.sequence)
    ).scalars().all()
    grouped: dict[int, list] = {}
    for h in rows:
        grouped.setdefault(h.record_id, []).append(h)
    return grouped


def _lab_by_record(db: Session, record_ids: list[int]) -> dict[int, list]:
    rows = db.execute(
        select(LabResult)
        .where(LabResult.record_id.in_(record_ids))
        .order_by(LabResult.record_id, LabResult.result_id)
    ).scalars().all()
    grouped: dict[int, list] = {}
    for r in rows:
        grouped.setdefault(r.record_id, []).append(r)
    return grouped


def _exam_by_record(db: Session, record_ids: list[int]) -> dict[int, list]:
    rows = db.execute(
        select(ExamReport)
        .where(ExamReport.record_id.in_(record_ids))
        .order_by(ExamReport.record_id, ExamReport.exam_id)
    ).scalars().all()
    grouped: dict[int, list] = {}
    for r in rows:
        grouped.setdefault(r.record_id, []).append(r)
    return grouped


def _add_lab_results(doc: Document, labs) -> None:
    if not labs:
        return
    doc.add_heading("检验结果", level=3)
    has_test = any((lab.test_name or "").strip() for lab in labs)
    cols = 6 if has_test else 5
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    headers = (["检验名称", "项目", "结果", "单位", "参考范围", "异常"]
               if has_test else ["项目", "结果", "单位", "参考范围", "异常"])
    for i, text in enumerate(headers):
        p = table.rows[0].cells[i].paragraphs[0]
        _style_run(p.add_run(text), bold=True)
    for lab in labs:
        cells = table.add_row().cells
        if has_test:
            values = [
                lab.test_name or "",
                lab.item_name,
                lab.result_value,
                lab.unit,
                lab.reference_range,
                ABNORMAL.get(lab.abnormal_flag or 0, ""),
            ]
        else:
            values = [
                lab.item_name,
                lab.result_value,
                lab.unit,
                lab.reference_range,
                ABNORMAL.get(lab.abnormal_flag or 0, ""),
            ]
        for i, text in enumerate(values):
            p = cells[i].paragraphs[0]
            _style_run(p.add_run(text or ""))


def _add_exams(doc: Document, exams) -> None:
    if not exams:
        return
    doc.add_heading("检查报告", level=3)
    for exam in exams:
        p = doc.add_paragraph()
        _style_run(p.add_run(exam.item_name), bold=True)
        _add_kv(doc, "所见：", exam.finding)
        _add_kv(doc, "结论：", exam.conclusion)


def generate_case_report(db: Session, course_id: str) -> BytesIO | None:
    records = db.execute(
        select(InfoRecord)
        .where(InfoRecord.course_id == course_id, InfoRecord.is_deleted.is_(False))
        .order_by(InfoRecord.visit_no)
    ).scalars().all()
    if not records:
        return None

    patient = db.scalar(select(InfoPatient).where(
        InfoPatient.patient_id == records[0].patient_id))
    herbs_map = _herbs_by_record(db, [r.record_id for r in records])
    lab_map = _lab_by_record(db, [r.record_id for r in records])
    exam_map = _exam_by_record(db, [r.record_id for r in records])

    # 关联的学习心得与导师点评
    notes = db.execute(
        select(LearningNote)
        .where(LearningNote.record_id.in_([r.record_id for r in records]),
               LearningNote.is_deleted.is_(False))
        .order_by(LearningNote.created_at)
    ).scalars().all()
    comments_by_note: dict[str, list] = {}
    if notes:
        note_ids = [n.note_id for n in notes]
        comments = db.execute(
            select(MentorComment)
            .where(MentorComment.target_type == 0,
                   MentorComment.target_id.in_(note_ids))
            .order_by(MentorComment.commented_at)
        ).scalars().all()
        for c in comments:
            comments_by_note.setdefault(c.target_id, []).append(c)

    doc = Document()
    title = doc.add_heading("病案报告", level=0)
    _style_run(title.runs[0], size=22, bold=True)

    # 一、基本信息
    doc.add_heading("一、基本信息", level=1)
    _add_kv(doc, "患者：",
            f"{patient.patient_name}（{'男' if patient.gender else '女'}）"
            if patient else records[0].patient_id)
    _add_kv(doc, "就诊日期：",
            f"{records[0].clinic_date} 至 {records[-1].clinic_date}"
            f"（共 {len(records)} 诊）")
    _add_kv(doc, "带教老师：", _mentor_name(db, records[0].mentor_id))
    _add_kv(doc, "就诊科室：", records[0].department)

    # 二、病程记录
    doc.add_heading("二、病程记录", level=1)
    for rec in records:
        doc.add_heading(f"第 {rec.visit_no} 诊（{rec.clinic_date}）", level=2)
        n = db.get(CaseNarrative, rec.record_id)
        d = db.get(Diagnosis, rec.record_id)
        t = db.get(Treatment, rec.record_id)
        herbs = herbs_map.get(rec.record_id, [])

        if n is not None:
            _add_kv(doc, "主诉：", n.complaint)
            _add_kv(doc, "现病史：", n.present_illness)
            if rec.visit_no == 1:
                # 既往史、个人史、过敏史“只记一次”，仅首诊展示
                _add_kv(doc, "既往史：", n.past_history)
                _add_kv(doc, "个人史及婚育史：", n.personal_history)
                _add_kv(doc, "过敏史：", n.allergy_history)
            tongue = "，".join(x for x in (n.body_of_tongue, n.fur_of_tongue) if x)
            _add_kv(doc, "舌象：", tongue)
            _add_kv(doc, "脉象：", n.pulse)
            _add_kv(doc, "其他四诊：", n.other_cond)
            _add_kv(doc, "体格检查：", n.physical_exam)
            _add_kv(doc, "辅助检查：", n.auxiliary_exam)

        if d is not None:
            _add_kv(doc, "中医病名：", d.tcm_disease)
            _add_kv(doc, "证型：", d.syndrome)
            _add_kv(doc, "西医诊断：", d.wm_diagnosis)
            _add_kv(doc, "辨证分析：", d.patterns_analysis)
            _add_kv(doc, "鉴别诊断：", d.differential_diagnosis)

        if t is not None:
            _add_kv(doc, "治法：", t.treatment_principle)
            formula = t.formula_name
            if t.dose_count:
                formula = f"{formula}（{t.dose_count} 付）" if formula else f"{t.dose_count} 付"
            _add_kv(doc, "方剂：", formula)
            _add_kv(doc, "煎煮法：", t.decoction)
            _add_kv(doc, "用法：", t.usage)

        if herbs:
            _add_prescription(doc, herbs)

        if t is not None:
            _add_kv(doc, "西药：", t.western_medicine)
            _add_kv(doc, "医嘱：", t.advice)
            _add_kv(doc, "其他治疗：", t.other_treatment)

        _add_lab_results(doc, lab_map.get(rec.record_id, []))
        _add_exams(doc, exam_map.get(rec.record_id, []))

    # 三、学习心得
    if notes:
        doc.add_heading("三、学习心得", level=1)
        for note in notes:
            type_label = NOTE_TYPES.get(note.note_type or 0, "心得")
            doc.add_heading(f"{note.title}（{type_label}）", level=2)
            for line in _md_to_plain(note.content_md):
                _style_run(doc.add_paragraph().add_run(line))
            for c in comments_by_note.get(note.note_id, []):
                p = doc.add_paragraph()
                _style_run(p.add_run("导师点评："), bold=True)
                _style_run(p.add_run(c.content))

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _mentor_name(db: Session, mentor_id: str | None) -> str:
    if not mentor_id:
        return ""
    from ..models import Mentor
    mentor = db.get(Mentor, mentor_id)
    return mentor.mentor_name if mentor else ""


def generate_notes_report(db: Session, note_ids: list[str]) -> BytesIO | None:
    """把选中的学习心得导出为 Word 汇编（含导师点评）。"""
    notes = db.execute(
        select(LearningNote)
        .where(LearningNote.note_id.in_(note_ids),
               LearningNote.is_deleted.is_(False))
        .order_by(LearningNote.created_at)
    ).scalars().all()
    if not notes:
        return None

    comments_by_note: dict[str, list] = {}
    comments = db.execute(
        select(MentorComment)
        .where(MentorComment.target_type == 0,
               MentorComment.target_id.in_(note_ids))
        .order_by(MentorComment.commented_at)
    ).scalars().all()
    for c in comments:
        comments_by_note.setdefault(c.target_id, []).append(c)

    doc = Document()
    title = doc.add_heading("跟师学习心得汇编", level=0)
    _style_run(title.runs[0], size=22, bold=True)

    for note in notes:
        type_label = NOTE_TYPES.get(note.note_type or 0, "心得")
        doc.add_heading(f"{note.title}（{type_label}）", level=1)
        _add_kv(doc, "记录时间：", note.created_at.strftime("%Y-%m-%d"))
        _add_kv(doc, "带教老师：", _mentor_name(db, note.mentor_id))
        for line in _md_to_plain(note.content_md):
            _style_run(doc.add_paragraph().add_run(line))
        for c in comments_by_note.get(note.note_id, []):
            p = doc.add_paragraph()
            _style_run(p.add_run("导师点评："), bold=True)
            _style_run(p.add_run(c.content))

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
